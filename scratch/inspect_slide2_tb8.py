import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide2 = prs.slides[1]
tb8 = None
for s in slide2.shapes:
    if s.name == 'TextBox 8':
        tb8 = s
        break

if tb8:
    for i, p in enumerate(tb8.text_frame.paragraphs):
        print(f'P{i}: text={repr(p.text)}')
        for r in p.runs:
            c = None
            try:
                c = r.font.color.rgb
            except:
                pass
            print(f'   run text={repr(r.text)}, font={r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.font.bold}, underline={r.font.underline}, color={c}')
