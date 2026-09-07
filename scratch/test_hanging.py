import pptx
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

def format_paragraph(p, level=0):
    pPr = p._p.get_or_add_pPr()
    for c in list(pPr):
        if any(c.tag.endswith(x) for x in ['buChar', 'buFont', 'buAutoNum', 'buSzPct', 'buSzPts', 'buNone']):
            pPr.remove(c)
    pPr.append(OxmlElement('a:buNone'))
    
    if level == 0:
        # Main section header (❖)
        pPr.set('marL', '0')
        pPr.set('indent', '0')
    elif level == 1:
        # Main pointer (•)
        # hanging indent: left margin 24pt (304800 EMU), indent -16pt (-203200 EMU)
        pPr.set('marL', '304800')
        pPr.set('indent', '-203200')
    elif level == 2:
        # Sub-bullet (–)
        # hanging indent: left margin 50pt (635000 EMU), indent -16pt (-203200 EMU)
        pPr.set('marL', '635000')
        pPr.set('indent', '-203200')

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide2 = prs.slides[1]
tb8 = [s for s in slide2.shapes if s.name == 'TextBox 8'][0]
tb8.text_frame.clear()

p0 = tb8.text_frame.paragraphs[0]
format_paragraph(p0, level=0)
r0 = p0.add_run()
r0.text = '❖  Proposed Solution (Describe your Idea/Solution/Prototype)'
r0.font.name = 'Arial'
r0.font.size = Pt(19)
r0.font.bold = True
r0.font.underline = True
r0.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

p1 = tb8.text_frame.add_paragraph()
format_paragraph(p1, level=1)
p1.space_before = Pt(8)
r1_b = p1.add_run()
r1_b.text = '•  '
r1_b.font.name = 'Arial'
r1_b.font.size = Pt(15)
r1_b.font.bold = True
r1_b.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
r1_t = p1.add_run()
r1_t.text = 'Detailed Explanation of the Proposed Solution'
r1_t.font.name = 'Arial'
r1_t.font.size = Pt(15)
r1_t.font.bold = True
r1_t.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

p2 = tb8.text_frame.add_paragraph()
format_paragraph(p2, level=2)
p2.space_before = Pt(3)
r2_d = p2.add_run()
r2_d.text = '–  '
r2_d.font.name = 'Arial'
r2_d.font.size = Pt(13)
r2_d.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
r2_p = p2.add_run()
r2_p.text = 'Platform Overview: '
r2_p.font.name = 'Arial'
r2_p.font.size = Pt(13)
r2_p.font.bold = True
r2_p.font.color.rgb = RGBColor(0x21, 0x25, 0x29)
r2_t = p2.add_run()
r2_t.text = 'AI-powered mobile/web platform translating primary school content into tribal & regional mother tongues in real time with synchronized text and audio.'
r2_t.font.name = 'Arial'
r2_t.font.size = Pt(13)
r2_t.font.bold = False
r2_t.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

prs.save('scratch/test_hanging.pptx')
print('Saved test_hanging.pptx')
