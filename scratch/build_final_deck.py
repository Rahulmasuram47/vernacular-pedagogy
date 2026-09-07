import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
import shutil, os

# Template source
template_pptx = r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx'
prs = pptx.Presentation(template_pptx)

# 1. Delete Slide 7 ("Important Instructions")
rId = prs.slides._sldIdLst[6].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[6]
assert len(prs.slides) == 6, f'Expected 6 slides, got {len(prs.slides)}'
print('Verified: Slide 7 deleted. Total slide count is 6.')

# Helper to format paragraph with custom bullets, margins, and hanging indents
def format_p(p, level=0, space_before=2, space_after=2, line_spacing=1.12):
    pPr = p._p.get_or_add_pPr()
    # Remove any existing bullet definitions to prevent double bullets
    for c in list(pPr):
        if any(c.tag.endswith(x) for x in ['buChar', 'buFont', 'buAutoNum', 'buSzPct', 'buSzPts', 'buNone']):
            pPr.remove(c)
    pPr.append(OxmlElement('a:buNone'))
    
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    
    if level == 0:
        # Top heading (❖)
        pPr.set('marL', '0')
        pPr.set('indent', '0')
    elif level == 1:
        # Main Pointer (•) -> marL=24pt (304800 EMU), indent=-16pt (-203200 EMU)
        pPr.set('marL', '304800')
        pPr.set('indent', '-203200')
    elif level == 2:
        # Sub-bullet (–) -> marL=50pt (635000 EMU), indent=-16pt (-203200 EMU)
        pPr.set('marL', '635000')
        pPr.set('indent', '-203200')

def add_pointer(tf, text, is_first=False, font_size=14.5, space_before=6):
    p = tf.paragraphs[0] if (is_first and len(tf.paragraphs) == 1 and not tf.paragraphs[0].text) else tf.add_paragraph()
    format_p(p, level=1, space_before=space_before, space_after=2, line_spacing=1.15)
    
    r_dot = p.add_run()
    r_dot.text = '•  '
    r_dot.font.name = 'Arial'
    r_dot.font.size = Pt(font_size)
    r_dot.font.bold = True
    r_dot.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    r_txt = p.add_run()
    r_txt.text = text
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(font_size)
    r_txt.font.bold = True
    r_txt.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return p

def add_sub_bullet(tf, text, bold_prefix='', font_size=12.5, space_before=1.5):
    p = tf.add_paragraph()
    format_p(p, level=2, space_before=space_before, space_after=1.5, line_spacing=1.12)
    
    r_dash = p.add_run()
    r_dash.text = '–  '
    r_dash.font.name = 'Arial'
    r_dash.font.size = Pt(font_size)
    r_dash.font.bold = True
    r_dash.font.color.rgb = RGBColor(0x5A, 0x6A, 0x85)
    
    if bold_prefix:
        r_pfx = p.add_run()
        r_pfx.text = bold_prefix + ': '
        r_pfx.font.name = 'Arial'
        r_pfx.font.size = Pt(font_size)
        r_pfx.font.bold = True
        r_pfx.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)
        
    r_txt = p.add_run()
    r_txt.text = text
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(font_size)
    r_txt.font.bold = False
    r_txt.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return p

# Update Team Name badge on slides 2 to 6
def update_team_oval(slide):
    for shape in slide.shapes:
        if 'Oval' in shape.name or (shape.has_text_frame and 'Your' in shape.text_frame.text and 'Team' in shape.text_frame.text):
            shape.text_frame.clear()
            p1 = shape.text_frame.paragraphs[0]
            format_p(p1, level=0, space_before=0, space_after=0)
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = 'Team'
            r1.font.name = 'Calibri'
            r1.font.size = Pt(17)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            
            p2 = shape.text_frame.add_paragraph()
            format_p(p2, level=0, space_before=0, space_after=0)
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = 'Bytes'
            r2.font.name = 'Calibri'
            r2.font.size = Pt(17)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            break

for i in range(1, 6):
    update_team_oval(prs.slides[i])

# ==========================================
# SLIDE 1 — TITLE PAGE
# ==========================================
slide1 = prs.slides[0]
tb9 = [s for s in slide1.shapes if s.name == 'TextBox 9'][0]
tb9.text_frame.clear()
tb9.text_frame.word_wrap = True
tb9.left = Emu(450000)
tb9.top = Emu(2050000)
tb9.width = Emu(6300000)
tb9.height = Emu(4500000)

fields = [
    ('Problem Statement ID –', ' SIH26042 / 042'),
    ('Problem Statement Title –', ' AI-Powered Vernacular Pedagogy and Real-Time Translation Tool for Mother Tongue-Based Primary Education'),
    ('Theme –', ' Education / EdTech'),
    ('PS Category –', ' Software'),
    ('Team ID –', ' [ENTER YOUR OFFICIAL SIH TEAM ID]'),
    ('Team Name (Registered on portal) –', ' Team Bytes')
]

for idx, (lbl, val) in enumerate(fields):
    p = tb9.text_frame.paragraphs[0] if idx == 0 else tb9.text_frame.add_paragraph()
    format_p(p, level=1, space_before=4, space_after=10, line_spacing=1.15)
    
    r_b = p.add_run()
    r_b.text = '•  '
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(16.5)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    r_l = p.add_run()
    r_l.text = lbl
    r_l.font.name = 'Arial'
    r_l.font.size = Pt(16.5)
    r_l.font.bold = True
    r_l.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    
    r_v = p.add_run()
    r_v.text = val
    r_v.font.name = 'Arial'
    r_v.font.size = Pt(15.5)
    r_v.font.bold = False
    if 'ENTER YOUR OFFICIAL' in val:
        r_v.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r_v.font.bold = True
    elif 'Team Bytes' in val or 'SIH26042' in val:
        r_v.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        r_v.font.bold = True
    else:
        r_v.font.color.rgb = RGBColor(0x21, 0x25, 0x29)

# ==========================================
# SLIDE 2 — IDEA TITLE / PROPOSED SOLUTION
# ==========================================
slide2 = prs.slides[1]
title2 = [s for s in slide2.shapes if s.name == 'Title 1'][0]
title2.text_frame.clear()
p_t2 = title2.text_frame.paragraphs[0]
format_p(p_t2, level=0, space_before=0, space_after=0)
p_t2.alignment = PP_ALIGN.CENTER
r_t2 = p_t2.add_run()
r_t2.text = 'IDEA TITLE: VERNACULAR VOICE'
r_t2.font.name = 'Times New Roman'
r_t2.font.size = Pt(32)
r_t2.font.bold = True
r_t2.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s2 = [s for s in slide2.shapes if s.name == 'TextBox 8'][0]
tb8_s2.left = Emu(450000)
tb8_s2.top = Emu(1300000)
tb8_s2.width = Emu(11300000)
tb8_s2.height = Emu(4900000)
tb8_s2.text_frame.clear()
tb8_s2.text_frame.word_wrap = True

# Main Section Header
p_head2 = tb8_s2.text_frame.paragraphs[0]
format_p(p_head2, level=0, space_before=0, space_after=5)
r_sym2 = p_head2.add_run()
r_sym2.text = '❖  '
r_sym2.font.name = 'Arial'
r_sym2.font.size = Pt(18.5)
r_sym2.font.bold = True
r_sym2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

r_txt2 = p_head2.add_run()
r_txt2.text = 'Proposed Solution (Describe your Idea/Solution/Prototype)'
r_txt2.font.name = 'Arial'
r_txt2.font.size = Pt(18.5)
r_txt2.font.bold = True
r_txt2.font.underline = True
r_txt2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

add_pointer(tb8_s2.text_frame, 'Detailed Explanation of the Proposed Solution', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s2.text_frame, 'AI-powered mobile/web platform that translates and converts primary-school educational content into tribal and regional mother tongues in real time.', 'Platform Overview', font_size=12)
add_sub_bullet(tb8_s2.text_frame, 'Combines speech-to-text (STT), text-to-speech (TTS), and neural translation to let teachers and students interact intuitively in their native language.', 'Multimodal Core', font_size=12)
add_sub_bullet(tb8_s2.text_frame, 'Teacher/student inputs content (text or speech) in source language (e.g., English/Hindi); system translates to target mother tongue and synthesizes matching audio.', 'Workflow', font_size=12)
add_sub_bullet(tb8_s2.text_frame, 'Generates bilingual learning material (synchronized text + audio) with offline-first local caching so it functions reliably in low-connectivity classrooms.', 'Offline Caching', font_size=12)

add_pointer(tb8_s2.text_frame, 'How It Addresses the Problem', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s2.text_frame, 'Removes the language barrier between standardized state curricula and mother-tongue primary learners, advancing NEP 2020 mandates.', 'Curriculum Alignment', font_size=12)
add_sub_bullet(tb8_s2.text_frame, 'Gives teachers an instant, automated tool to localize lessons and worksheets instead of relying on scarce human translators.', 'Teacher Empowerment', font_size=12)

add_pointer(tb8_s2.text_frame, 'Innovation and Uniqueness of the Solution', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s2.text_frame, 'Modular language-plugin architecture allows new tribal and Indian regional languages to be onboarded without rebuilding the core system.', 'Modular Architecture', font_size=12)
add_sub_bullet(tb8_s2.text_frame, 'Quantized offline AI inference engineered specifically for entry-level, budget Android devices, unlike conventional cloud-only EdTech tools.', 'Edge-Optimized AI', font_size=12)

# ==========================================
# SLIDE 3 — TECHNICAL APPROACH
# ==========================================
slide3 = prs.slides[2]
title3 = [s for s in slide3.shapes if s.name == 'Title 1'][0]
title3.text_frame.clear()
p_t3 = title3.text_frame.paragraphs[0]
format_p(p_t3, level=0, space_before=0, space_after=0)
p_t3.alignment = PP_ALIGN.CENTER
r_t3 = p_t3.add_run()
r_t3.text = 'TECHNICAL APPROACH'
r_t3.font.name = 'Times New Roman'
r_t3.font.size = Pt(34)
r_t3.font.bold = True
r_t3.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s3 = [s for s in slide3.shapes if s.name == 'TextBox 8'][0]
tb8_s3.left = Emu(450000)
tb8_s3.top = Emu(1250000)
tb8_s3.width = Emu(11300000)
tb8_s3.height = Emu(3150000)
tb8_s3.text_frame.clear()
tb8_s3.text_frame.word_wrap = True

add_pointer(tb8_s3.text_frame, 'Technologies to be Used (Programming Languages, Frameworks, Hardware)', is_first=True, font_size=14.5, space_before=2)
add_sub_bullet(tb8_s3.text_frame, 'Flutter / React Native (cross-platform, lightweight client, highly responsive on low-end Android smartphones).', 'Client Application', font_size=12)
add_sub_bullet(tb8_s3.text_frame, 'Python (FastAPI microservices) for high-performance API orchestration, model serving, and asynchronous sync.', 'Backend Framework', font_size=12)
add_sub_bullet(tb8_s3.text_frame, 'IndicTrans2 / fine-tuned multilingual NMT for Indian & tribal languages; Whisper (offline-optimized) / Vosk for STT; Indic-TTS / Coqui TTS for speech synthesis.', 'AI & ML Models', font_size=12)
add_sub_bullet(tb8_s3.text_frame, 'TensorFlow Lite & ONNX Runtime for INT8 quantized offline execution; SQLite for local content caching; Firebase / PostgreSQL for optional cloud sync.', 'Edge Inference & DB', font_size=12)

add_pointer(tb8_s3.text_frame, 'Methodology and Process for Implementation (Flow Charts / Images / Working Prototype)', font_size=14.5, space_before=4)
add_sub_bullet(tb8_s3.text_frame, 'Step 1: Input (Text/Voice) → Step 2: STT Conversion → Step 3: NMT Translation to Mother Tongue → Step 4: TTS Native Audio Synthesis → Step 5: Bilingual Delivery (Text+Audio) & SQLite Cache.', 'End-to-End Pipeline', font_size=12)

# Flowchart image on Slide 3
flowchart_path = 'scratch/methodology_flowchart.png'
img_left = Emu(700000)
img_top = Emu(4500000)
img_width = Emu(10800000)
img_height = Emu(1700000)
slide3.shapes.add_picture(flowchart_path, img_left, img_top, width=img_width, height=img_height)

# ==========================================
# SLIDE 4 — FEASIBILITY AND VIABILITY
# ==========================================
slide4 = prs.slides[3]
title4 = [s for s in slide4.shapes if s.name == 'Title 1'][0]
title4.text_frame.clear()
p_t4 = title4.text_frame.paragraphs[0]
format_p(p_t4, level=0, space_before=0, space_after=0)
p_t4.alignment = PP_ALIGN.CENTER
r_t4 = p_t4.add_run()
r_t4.text = 'FEASIBILITY AND VIABILITY'
r_t4.font.name = 'Times New Roman'
r_t4.font.size = Pt(34)
r_t4.font.bold = True
r_t4.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s4 = [s for s in slide4.shapes if s.name == 'TextBox 8'][0]
tb8_s4.left = Emu(450000)
tb8_s4.top = Emu(1250000)
tb8_s4.width = Emu(11300000)
tb8_s4.height = Emu(4950000)
tb8_s4.text_frame.clear()
tb8_s4.text_frame.word_wrap = True

add_pointer(tb8_s4.text_frame, 'Analysis of the Feasibility of the Idea', is_first=True, font_size=14.5, space_before=2)
add_sub_bullet(tb8_s4.text_frame, 'Built on existing, proven AI/ML and speech technologies (IndicTrans2, Whisper, Indic-TTS) with verified Indian language benchmarks.', 'Proven Foundations', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Designed specifically for low-end Android smartphones prevalent in rural classrooms, ensuring zero prohibitive hardware barrier.', 'Hardware Compatibility', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Modular architecture allows easy, continuous addition of more Indian and tribal languages without platform refactoring.', 'Modular Scalability', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Offline-capable functionality reduces dependency on continuous internet connectivity and eliminates ongoing data subscription costs.', 'Offline Feasibility', font_size=12)

add_pointer(tb8_s4.text_frame, 'Potential Challenges and Risks', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s4.text_frame, 'Limited parallel datasets and acoustic corpora available for low-resource tribal languages (e.g., Santhali, Mundari, Ho).', 'Dataset Scarcity', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Accuracy of translation and speech recognition for regional dialects, non-standard accents, and phoneme variations.', 'Accuracy & Nuance', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Real-time response latency on low-RAM budget devices; reliable operation under severe connectivity constraints.', 'Latency & Resources', font_size=12)

add_pointer(tb8_s4.text_frame, 'Strategies for Overcoming These Challenges', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s4.text_frame, 'Use lightweight and INT8 quantized models optimized for on-device inference via TFLite and ONNX Runtime to ensure sub-second response.', 'Quantized Edge Runtime', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Pre-cache frequently used primary curriculum units and lessons locally, providing instant playback and graceful offline fallbacks.', 'Local Cache & Fallbacks', font_size=12)
add_sub_bullet(tb8_s4.text_frame, 'Continuously expand tribal language resources using validated educational datasets and community-in-the-loop teacher validation.', 'Continuous Expansion', font_size=12)

# ==========================================
# SLIDE 5 — IMPACT AND BENEFITS
# ==========================================
slide5 = prs.slides[4]
title5 = [s for s in slide5.shapes if s.name == 'Title 1'][0]
title5.text_frame.clear()
p_t5 = title5.text_frame.paragraphs[0]
format_p(p_t5, level=0, space_before=0, space_after=0)
p_t5.alignment = PP_ALIGN.CENTER
r_t5 = p_t5.add_run()
r_t5.text = 'IMPACT AND BENEFITS'
r_t5.font.name = 'Times New Roman'
r_t5.font.size = Pt(34)
r_t5.font.bold = True
r_t5.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s5 = [s for s in slide5.shapes if s.name == 'TextBox 8'][0]
tb8_s5.left = Emu(450000)
tb8_s5.top = Emu(1250000)
tb8_s5.width = Emu(11300000)
tb8_s5.height = Emu(4950000)
tb8_s5.text_frame.clear()
tb8_s5.text_frame.word_wrap = True

add_pointer(tb8_s5.text_frame, 'Potential Impact on the Target Audience', is_first=True, font_size=14.5, space_before=2)
add_sub_bullet(tb8_s5.text_frame, 'Enables mother-tongue-based primary education as envisioned by NEP 2020, eliminating early childhood cognitive exclusion.', 'NEP 2020 Vision', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Makes educational content fully accessible to tribal-language-speaking children, drastically reducing learning poverty and dropout rates.', 'Equitable Access', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Bridges comprehension gaps between standardized state textbooks and local spoken dialects, boosting foundational numeracy and literacy.', 'Learning Outcomes', font_size=12)

add_pointer(tb8_s5.text_frame, 'Benefits of the Solution (Social, Economic, Environmental, etc.)', font_size=14.5, space_before=5)
add_sub_bullet(tb8_s5.text_frame, 'Helps teachers create bilingual learning materials instantly, saving hundreds of hours of manual lesson preparation.', 'Teacher Empowerment', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Supports voice-based learning for young, pre-literate learners, making foundational education engaging, confident, and intuitive.', 'Multimodal Inclusivity', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Reduces language barriers between teachers and students, creating an interactive, culturally affirming classroom environment.', 'Classroom Harmony', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Scalable to additional Indian regional and tribal languages without expensive physical infrastructure investments.', 'High Scalability', font_size=12)
add_sub_bullet(tb8_s5.text_frame, 'Offline capability ensures 100% usability and equity in low-connectivity rural and remote tribal schools.', 'Zero-Connectivity Reach', font_size=12)

# ==========================================
# SLIDE 6 — RESEARCH AND REFERENCES
# ==========================================
slide6 = prs.slides[5]
title6 = [s for s in slide6.shapes if s.name == 'Title 1'][0]
title6.text_frame.clear()
p_t6 = title6.text_frame.paragraphs[0]
format_p(p_t6, level=0, space_before=0, space_after=0)
p_t6.alignment = PP_ALIGN.CENTER
r_t6 = p_t6.add_run()
r_t6.text = 'RESEARCH AND REFERENCES'
r_t6.font.name = 'Times New Roman'
r_t6.font.size = Pt(34)
r_t6.font.bold = True
r_t6.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s6 = [s for s in slide6.shapes if s.name == 'TextBox 8'][0]
tb8_s6.left = Emu(450000)
tb8_s6.top = Emu(1250000)
tb8_s6.width = Emu(11300000)
tb8_s6.height = Emu(4950000)
tb8_s6.text_frame.clear()
tb8_s6.text_frame.word_wrap = True

add_pointer(tb8_s6.text_frame, 'Details / Links of the Reference and Research Work', is_first=True, font_size=14.5, space_before=2)
add_sub_bullet(tb8_s6.text_frame, 'Government of Jharkhand — Smart India Hackathon 2026 Problem Statement ID: SIH26042 / 042.', 'Problem Statement', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'Official Smart India Hackathon 2026 Portal (https://sih.gov.in) — Guidelines, Submission Template & EdTech Domain Specs.', 'SIH 2026 Portal', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'IndicTrans2: Towards High-Quality & Accessible Machine Translation for Indian Languages (AI4Bharat / Gala et al., 2023).', 'Neural Translation Model', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'OpenAI Whisper: Robust Speech Recognition via Large-Scale Weak Supervision (Radford et al.) & AlphaCephei Vosk Offline Embedded Engine.', 'Speech-to-Text Research', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'Indic-TTS: Multilingual Text-to-Speech Synthesis System for Indian Languages (IIT Madras) & Coqui Open-Source Deep Learning TTS.', 'Text-to-Speech Technology', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'Bhashini (National Language Translation Mission, MeitY) & AI4Bharat Open Linguistic Corpora for Indian & Tribal Languages.', 'Linguistic Datasets', font_size=12)
add_sub_bullet(tb8_s6.text_frame, 'National Education Policy (NEP) 2020, Ministry of Education, Govt. of India — Guidelines on Multilingual & Mother-Tongue Pedagogy.', 'Policy Framework', font_size=12)

# Save target PPTX
output_pptx = r'scratch\SIH26042_TeamBytes_IdeaPresentation.pptx'
prs.save(output_pptx)
print('Successfully generated final PPTX at:', output_pptx)
