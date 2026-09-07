import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')

# 1. Delete Slide 7 ("Important Instructions")
rId = prs.slides._sldIdLst[6].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[6]
print('Slide count after deleting Slide 7:', len(prs.slides))

# Function to update Team Name oval
def update_team_oval(slide):
    for shape in slide.shapes:
        if 'Oval' in shape.name or (shape.has_text_frame and 'Your' in shape.text_frame.text and 'Team' in shape.text_frame.text):
            shape.text_frame.clear()
            p1 = shape.text_frame.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = 'Team'
            r1.font.name = 'Calibri'
            r1.font.size = Pt(17)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            
            p2 = shape.text_frame.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = 'Bytes'
            r2.font.name = 'Calibri'
            r2.font.size = Pt(17)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            break

# Update ovals on slides 2 to 6
for i in range(1, 6):
    update_team_oval(prs.slides[i])

# Helper function to add a main pointer
def add_pointer(tf, text, is_first=False, font_size=15):
    p = tf.paragraphs[0] if (is_first and len(tf.paragraphs) == 1 and not tf.paragraphs[0].text) else tf.add_paragraph()
    p.space_before = Pt(6)
    p.space_after = Pt(2)
    p.line_spacing = 1.15
    
    r_dot = p.add_run()
    r_dot.text = '•  '
    r_dot.font.name = 'Arial'
    r_dot.font.size = Pt(font_size)
    r_dot.font.bold = True
    r_dot.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    r_txt = p.add_run()
    r_txt.text = text
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(font_size)
    r_txt.font.bold = True
    r_txt.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    return p

# Helper function to add sub-bullets
def add_sub_bullet(tf, text, bold_prefix='', font_size=13):
    p = tf.add_paragraph()
    p.level = 1
    p.space_before = Pt(1.5)
    p.space_after = Pt(2)
    p.line_spacing = 1.12
    
    r_dash = p.add_run()
    r_dash.text = '    –  '
    r_dash.font.name = 'Arial'
    r_dash.font.size = Pt(font_size)
    r_dash.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    
    if bold_prefix:
        r_pfx = p.add_run()
        r_pfx.text = bold_prefix + ': '
        r_pfx.font.name = 'Arial'
        r_pfx.font.size = Pt(font_size)
        r_pfx.font.bold = True
        r_pfx.font.color.rgb = RGBColor(0x21, 0x25, 0x29)
        
    r_txt = p.add_run()
    r_txt.text = text
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(font_size)
    r_txt.font.bold = False
    r_txt.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return p

# ==========================================
# SLIDE 1 — TITLE PAGE
# ==========================================
slide1 = prs.slides[0]
tb9 = [s for s in slide1.shapes if s.name == 'TextBox 9'][0]
tb9.text_frame.clear()
tb9.text_frame.word_wrap = True
tb9.left = Emu(450000)
tb9.top = Emu(2000000)
tb9.width = Emu(6200000)
tb9.height = Emu(4600000)

fields = [
    ('Problem Statement ID –', ' SIH26042 / 042'),
    ('Problem Statement Title –', ' AI-Powered Vernacular Pedagogy and Real-Time Translation Tool for Mother Tongue-Based Primary Education'),
    ('Theme –', ' Education / EdTech'),
    ('PS Category –', ' Software'),
    ('Team ID –', ' [ENTER YOUR OFFICIAL SIH TEAM ID]'),
    ('Team Name (Registered on portal) –', ' Team Bytes')
]

for idx, (lbl, val) in enumerate(fields):
    p = tb9.text_frame.paragraphs[0] if idx == 0 else tb9.text_frame.add_paragraph()
    p.space_after = Pt(14)
    p.line_spacing = 1.15
    
    r_b = p.add_run()
    r_b.text = '•  '
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(17)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    r_l = p.add_run()
    r_l.text = lbl
    r_l.font.name = 'Arial'
    r_l.font.size = Pt(17)
    r_l.font.bold = True
    r_l.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    
    r_v = p.add_run()
    r_v.text = val
    r_v.font.name = 'Arial'
    r_v.font.size = Pt(16)
    r_v.font.bold = False
    if 'ENTER YOUR OFFICIAL' in val:
        r_v.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r_v.font.bold = True
    elif 'Team Bytes' in val or 'SIH26042' in val:
        r_v.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        r_v.font.bold = True
    else:
        r_v.font.color.rgb = RGBColor(0x21, 0x25, 0x29)

# ==========================================
# SLIDE 2 — IDEA TITLE / PROPOSED SOLUTION
# ==========================================
slide2 = prs.slides[1]
title2 = [s for s in slide2.shapes if s.name == 'Title 1'][0]
title2.text_frame.clear()
p_t2 = title2.text_frame.paragraphs[0]
r_t2 = p_t2.add_run()
r_t2.text = 'IDEA TITLE: VERNACULAR VOICE'
r_t2.font.name = 'Times New Roman'
r_t2.font.size = Pt(32)
r_t2.font.bold = True
r_t2.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s2 = [s for s in slide2.shapes if s.name == 'TextBox 8'][0]
tb8_s2.left = Emu(450000)
tb8_s2.top = Emu(1350000)
tb8_s2.width = Emu(11300000)
tb8_s2.height = Emu(4850000)
tb8_s2.text_frame.clear()
tb8_s2.text_frame.word_wrap = True

# Underlined section header
p_head2 = tb8_s2.text_frame.paragraphs[0]
p_head2.space_before = Pt(0)
p_head2.space_after = Pt(7)
r_sym2 = p_head2.add_run()
r_sym2.text = '❖  '
r_sym2.font.name = 'Arial'
r_sym2.font.size = Pt(19)
r_sym2.font.bold = True
r_sym2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

r_txt2 = p_head2.add_run()
r_txt2.text = 'Proposed Solution (Describe your Idea/Solution/Prototype)'
r_txt2.font.name = 'Arial'
r_txt2.font.size = Pt(19)
r_txt2.font.bold = True
r_txt2.font.underline = True
r_txt2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

add_pointer(tb8_s2.text_frame, 'Detailed Explanation of the Proposed Solution', font_size=15)
add_sub_bullet(tb8_s2.text_frame, 'AI-powered mobile/web platform translating primary school content into tribal & regional mother tongues in real time.', 'Platform Overview', font_size=13)
add_sub_bullet(tb8_s2.text_frame, 'Combines speech-to-text (STT), neural machine translation (NMT), and text-to-speech (TTS) for natural two-way native interaction.', 'Multimodal Core', font_size=13)
add_sub_bullet(tb8_s2.text_frame, 'Generates synchronized bilingual learning materials (text + audio) with offline-first caching for zero-connectivity classrooms.', 'Offline Flow', font_size=13)

add_pointer(tb8_s2.text_frame, 'How It Addresses the Problem', font_size=15)
add_sub_bullet(tb8_s2.text_frame, 'Removes language barriers between standardized curricula and mother-tongue learners, fulfilling NEP 2020 primary mandates.', 'Bridging Curriculum Gap', font_size=13)
add_sub_bullet(tb8_s2.text_frame, 'Provides rural educators with an instant localization tool, eliminating complete reliance on scarce human native translators.', 'Teacher Empowerment', font_size=13)

add_pointer(tb8_s2.text_frame, 'Innovation and Uniqueness of the Solution', font_size=15)
add_sub_bullet(tb8_s2.text_frame, 'Modular language-plugin architecture allows rapid onboarding of new tribal/regional dialects without rebuilding the core engine.', 'Modular Architecture', font_size=13)
add_sub_bullet(tb8_s2.text_frame, 'Quantized offline on-device AI inference engineered specifically for ultra-low-cost Android devices, unlike cloud-bound tools.', 'Edge-Optimized AI', font_size=13)

# ==========================================
# SLIDE 3 — TECHNICAL APPROACH
# ==========================================
slide3 = prs.slides[2]
title3 = [s for s in slide3.shapes if s.name == 'Title 1'][0]
title3.text_frame.clear()
p_t3 = title3.text_frame.paragraphs[0]
r_t3 = p_t3.add_run()
r_t3.text = 'TECHNICAL APPROACH'
r_t3.font.name = 'Times New Roman'
r_t3.font.size = Pt(36)
r_t3.font.bold = True
r_t3.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s3 = [s for s in slide3.shapes if s.name == 'TextBox 8'][0]
tb8_s3.left = Emu(450000)
tb8_s3.top = Emu(1250000)
tb8_s3.width = Emu(11300000)
tb8_s3.height = Emu(3200000)
tb8_s3.text_frame.clear()
tb8_s3.text_frame.word_wrap = True

add_pointer(tb8_s3.text_frame, 'Technologies to be Used (Programming Languages, Frameworks, Hardware)', is_first=True, font_size=14.5)
add_sub_bullet(tb8_s3.text_frame, 'Flutter / React Native (lightweight, cross-platform, optimized for low-end Android smartphones).', 'Client Application', font_size=12.5)
add_sub_bullet(tb8_s3.text_frame, 'Python (FastAPI microservices) for high-efficiency model serving, API orchestration, and sync.', 'Backend Framework', font_size=12.5)
add_sub_bullet(tb8_s3.text_frame, 'IndicTrans2 (Multilingual NMT), OpenAI Whisper & Vosk (offline STT), Indic-TTS & Coqui TTS (native audio).', 'Core AI Models', font_size=12.5)
add_sub_bullet(tb8_s3.text_frame, 'TensorFlow Lite / ONNX Runtime for INT8 quantized offline execution; SQLite (offline cache) & PostgreSQL.', 'Edge & Storage', font_size=12.5)

add_pointer(tb8_s3.text_frame, 'Methodology and Process for Implementation (Flow Charts / Images / Working Prototype)', font_size=14.5)
add_sub_bullet(tb8_s3.text_frame, 'Step 1 (Input Content) → Step 2 (Speech-to-Text) → Step 3 (Neural Translation) → Step 4 (TTS Audio Generation) → Step 5 (Bilingual Delivery & Offline Caching).', 'End-to-End Pipeline', font_size=12.5)

# Add Flowchart Image to Slide 3
flowchart_path = 'scratch/methodology_flowchart.png'
img_left = Emu(750000)
img_top = Emu(4450000)
img_width = Emu(10700000)
img_height = Emu(1800000)
slide3.shapes.add_picture(flowchart_path, img_left, img_top, width=img_width, height=img_height)

# ==========================================
# SLIDE 4 — FEASIBILITY AND VIABILITY
# ==========================================
slide4 = prs.slides[3]
title4 = [s for s in slide4.shapes if s.name == 'Title 1'][0]
title4.text_frame.clear()
p_t4 = title4.text_frame.paragraphs[0]
r_t4 = p_t4.add_run()
r_t4.text = 'FEASIBILITY AND VIABILITY'
r_t4.font.name = 'Times New Roman'
r_t4.font.size = Pt(36)
r_t4.font.bold = True
r_t4.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s4 = [s for s in slide4.shapes if s.name == 'TextBox 8'][0]
tb8_s4.left = Emu(450000)
tb8_s4.top = Emu(1250000)
tb8_s4.width = Emu(11300000)
tb8_s4.height = Emu(4950000)
tb8_s4.text_frame.clear()
tb8_s4.text_frame.word_wrap = True

add_pointer(tb8_s4.text_frame, 'Analysis of the Feasibility of the Idea', is_first=True, font_size=15)
add_sub_bullet(tb8_s4.text_frame, 'Built on proven AI/ML models (IndicTrans2, Whisper, Indic-TTS) with verified benchmarks across Indian languages.', 'Proven AI Foundations', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Engineered specifically for low-end Android smartphones prevalent in rural and tribal schooling environments.', 'Hardware Compatibility', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Modular language-plugin architecture allows rapid expansion to new Indian and tribal dialects without system refactoring.', 'Modular Scalability', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Offline-first functionality eliminates ongoing internet data costs and operates reliably in remote shadow regions.', 'Zero-Cost Offline Flow', font_size=13)

add_pointer(tb8_s4.text_frame, 'Potential Challenges and Risks', font_size=15)
add_sub_bullet(tb8_s4.text_frame, 'Limited parallel datasets and acoustic corpora available for low-resource tribal languages (e.g. Santhali, Mundari, Ho).', 'Dataset Scarcity', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Dialectical variations and pronunciation nuances affecting real-time speech recognition and translation accuracy.', 'Linguistic Diversity', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Real-time inference latency and RAM/storage constraints on entry-level Android devices.', 'Hardware Latency', font_size=13)

add_pointer(tb8_s4.text_frame, 'Strategies for Overcoming These Challenges', font_size=15)
add_sub_bullet(tb8_s4.text_frame, 'Deploy INT8 quantized TFLite / ONNX models to achieve sub-second on-device inference with minimal memory footprint.', 'Quantized Edge Runtime', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Locally pre-cache frequently used primary curriculum units and lessons for instant zero-latency classroom playback.', 'Intelligent Local Cache', font_size=13)
add_sub_bullet(tb8_s4.text_frame, 'Continuously expand tribal language resources via validated educational corpora and community-in-the-loop feedback.', 'Community Data Expansion', font_size=13)

# ==========================================
# SLIDE 5 — IMPACT AND BENEFITS
# ==========================================
slide5 = prs.slides[4]
title5 = [s for s in slide5.shapes if s.name == 'Title 1'][0]
title5.text_frame.clear()
p_t5 = title5.text_frame.paragraphs[0]
r_t5 = p_t5.add_run()
r_t5.text = 'IMPACT AND BENEFITS'
r_t5.font.name = 'Times New Roman'
r_t5.font.size = Pt(36)
r_t5.font.bold = True
r_t5.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s5 = [s for s in slide5.shapes if s.name == 'TextBox 8'][0]
tb8_s5.left = Emu(450000)
tb8_s5.top = Emu(1250000)
tb8_s5.width = Emu(11300000)
tb8_s5.height = Emu(4950000)
tb8_s5.text_frame.clear()
tb8_s5.text_frame.word_wrap = True

add_pointer(tb8_s5.text_frame, 'Potential Impact on the Target Audience', is_first=True, font_size=15)
add_sub_bullet(tb8_s5.text_frame, 'Directly operationalizes the National Education Policy (NEP 2020) mandate for mother-tongue foundational schooling.', 'NEP 2020 Alignment', font_size=13)
add_sub_bullet(tb8_s5.text_frame, 'Prevents early childhood cognitive exclusion and learning dropouts among tribal and vernacular-speaking children.', 'Inclusive Access', font_size=13)
add_sub_bullet(tb8_s5.text_frame, 'Bridges the comprehension gap between state-level textbooks and local dialects, accelerating primary literacy and numeracy.', 'Enhanced Learning Gains', font_size=13)

add_pointer(tb8_s5.text_frame, 'Benefits of the Solution (Social, Economic, Environmental, etc.)', font_size=15)
add_sub_bullet(tb8_s5.text_frame, 'Empowers rural primary teachers to instantly generate interactive bilingual text and audio learning materials.', 'Teacher Assistance', font_size=13)
add_sub_bullet(tb8_s5.text_frame, 'Voice-first multimodal capability enables non-literate and early-grade learners to participate and comprehend intuitively.', 'Voice-Based Inclusivity', font_size=13)
add_sub_bullet(tb8_s5.text_frame, 'Eliminates language friction between teachers and pupils, fostering interactive, confident classroom environments.', 'Classroom Engagement', font_size=13)
add_sub_bullet(tb8_s5.text_frame, 'Readily scalable across additional regional and indigenous dialects without costly infrastructural overhaul.', 'Linguistic Scalability', font_size=13)
add_sub_bullet(tb8_s5.text_frame, '100% offline capability ensures uninterrupted, equitable learning in remote rural tribal schools with zero connectivity.', 'Rural Resilience', font_size=13)

# ==========================================
# SLIDE 6 — RESEARCH AND REFERENCES
# ==========================================
slide6 = prs.slides[5]
title6 = [s for s in slide6.shapes if s.name == 'Title 1'][0]
title6.text_frame.clear()
p_t6 = title6.text_frame.paragraphs[0]
r_t6 = p_t6.add_run()
r_t6.text = 'RESEARCH AND REFERENCES'
r_t6.font.name = 'Times New Roman'
r_t6.font.size = Pt(36)
r_t6.font.bold = True
r_t6.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

tb8_s6 = [s for s in slide6.shapes if s.name == 'TextBox 8'][0]
tb8_s6.left = Emu(450000)
tb8_s6.top = Emu(1250000)
tb8_s6.width = Emu(11300000)
tb8_s6.height = Emu(4950000)
tb8_s6.text_frame.clear()
tb8_s6.text_frame.word_wrap = True

add_pointer(tb8_s6.text_frame, 'Details / Links of the Reference and Research Work', is_first=True, font_size=15)
add_sub_bullet(tb8_s6.text_frame, 'Government of Jharkhand — Smart India Hackathon 2026 Problem Statement ID: SIH26042 / 042.', 'Problem Statement', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'Official Smart India Hackathon 2026 Portal (https://sih.gov.in) — Guidelines & EdTech Domain Specifications.', 'SIH 2026 Platform', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'AI4Bharat IndicTrans2: Towards High-Quality & Accessible Machine Translation for Indian Languages (Gala et al., 2023).', 'Neural Translation', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'OpenAI Whisper (Robust Speech Recognition via Weak Supervision) & AlphaCephei Vosk Offline Embedded STT.', 'Speech-to-Text Models', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'Indic-TTS (Multilingual TTS for Indian Languages, IIT Madras) & Coqui TTS Deep-Learning Audio Synthesis.', 'Text-to-Speech Engine', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'Bhashini (National Language Translation Mission, MeitY) & AI4Bharat Open Linguistic Corpora for Indian & Tribal Languages.', 'Tribal Datasets', font_size=13)
add_sub_bullet(tb8_s6.text_frame, 'National Education Policy (NEP 2020), Ministry of Education, Govt. of India — Multilingual & Mother-Tongue Pedagogy.', 'Policy Guidelines', font_size=13)

# Save final PPTX
output_pptx = r'scratch\SIH26042_TeamBytes_IdeaPresentation.pptx'
prs.save(output_pptx)
print('All 6 slides built and saved successfully to:', output_pptx)
