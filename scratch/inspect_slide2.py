import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide2 = prs.slides[1]
for s in slide2.shapes:
    print(f'Shape {s.name}: left={s.left}, top={s.top}, width={s.width}, height={s.height}')
    if s.has_text_frame:
        for i, p in enumerate(s.text_frame.paragraphs):
            print(f'   P{i}: text={repr(p.text)}')
