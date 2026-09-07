import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')

for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for shape in slide.shapes:
        info = f'Shape ID: {shape.shape_id}, Name: {shape.name}'
        print(info)
        if shape.has_text_frame:
            tf = shape.text_frame
            for p_idx, p in enumerate(tf.paragraphs):
                runs_info = []
                for r in p.runs:
                    fn = r.font.name
                    fs = r.font.size.pt if r.font.size else 'None'
                    b = r.font.bold
                    runs_info.append(f'\"{r.text}\" ({fn}, {fs}pt, bold={b})')
                print(f'    P{p_idx} [lvl={p.level}]: ' + ' '.join(runs_info))
