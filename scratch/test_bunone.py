import pptx
from pptx.oxml import OxmlElement

def remove_auto_bullets(p):
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        if any(child.tag.endswith(x) for x in ['buChar', 'buFont', 'buAutoNum', 'buSzPct', 'buSzPts', 'buNone']):
            pPr.remove(child)
    buNone = OxmlElement('a:buNone')
    pPr.append(buNone)
    # also reset indent / margin
    pPr.set('marL', '0')
    pPr.set('indent', '0')

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide3 = prs.slides[2]
tb = [s for s in slide3.shapes if s.name == 'TextBox 8'][0]
p0 = tb.text_frame.paragraphs[0]
remove_auto_bullets(p0)
print('Success removing bullets')
