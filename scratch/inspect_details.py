import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')

for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            print(f'-- Shape: {shape.name} (type={shape.shape_type})')
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                text = p.text
                if not text.strip():
                    continue
                font_names = set(r.font.name for r in p.runs if r.font.name)
                font_sizes = set(r.font.size.pt for r in p.runs if r.font.size)
                font_colors = set(str(r.font.color.rgb) for r in p.runs if r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB)
                bolds = set(r.font.bold for r in p.runs)
                print(f'   P{p_idx}: \"{text}\"')
                print(f'        fonts={font_names}, sizes={font_sizes}, bolds={bolds}, colors={font_colors}, align={p.alignment}')
