import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')

# 1. Delete Slide 7
rId = prs.slides._sldIdLst[6].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[6]
print('Slide count after deleting slide 7:', len(prs.slides))

# Function to update Team Name oval
def update_team_oval(slide):
    for shape in slide.shapes:
        if 'Oval' in shape.name or (shape.has_text_frame and 'Your' in shape.text_frame.text and 'Team' in shape.text_frame.text):
            shape.text_frame.clear()
            p1 = shape.text_frame.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            r1 = p1.add_run()
            r1.text = 'Team'
            r1.font.name = 'Calibri'
            r1.font.size = Pt(17)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            
            p2 = shape.text_frame.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = 'Bytes'
            r2.font.name = 'Calibri'
            r2.font.size = Pt(17)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            print(f'Updated oval on slide')
            break

# Update ovals on slides 2 to 6 (indices 1 to 5)
for i in range(1, 6):
    update_team_oval(prs.slides[i])

# ==================== SLIDE 1 ====================
slide1 = prs.slides[0]
tb9 = [s for s in slide1.shapes if s.name == 'TextBox 9'][0]
tb9.text_frame.clear()
tb9.text_frame.word_wrap = True

# Adjust position slightly if needed
tb9.left = Emu(450000)
tb9.top = Emu(2000000)
tb9.width = Emu(6200000)
tb9.height = Emu(4600000)

fields = [
    ('Problem Statement ID –', ' SIH26042 / 042'),
    ('Problem Statement Title –', ' AI-Powered Vernacular Pedagogy and Real-Time Translation Tool for Mother Tongue-Based Primary Education'),
    ('Theme –', ' Education / EdTech'),
    ('PS Category –', ' Software'),
    ('Team ID –', ' [ENTER YOUR OFFICIAL SIH TEAM ID]'),
    ('Team Name (Registered on portal) –', ' Team Bytes')
]

for idx, (lbl, val) in enumerate(fields):
    p = tb9.text_frame.paragraphs[0] if idx == 0 else tb9.text_frame.add_paragraph()
    p.space_after = Pt(14)
    p.line_spacing = 1.15
    
    r_bullet = p.add_run()
    r_bullet.text = '•  '
    r_bullet.font.name = 'Arial'
    r_bullet.font.size = Pt(17)
    r_bullet.font.bold = True
    r_bullet.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    r_lbl = p.add_run()
    r_lbl.text = lbl
    r_lbl.font.name = 'Arial'
    r_lbl.font.size = Pt(17)
    r_lbl.font.bold = True
    r_lbl.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    
    r_val = p.add_run()
    r_val.text = val
    r_val.font.name = 'Arial'
    r_val.font.size = Pt(16)
    r_val.font.bold = False
    if 'ENTER YOUR OFFICIAL' in val:
        r_val.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        r_val.font.bold = True
    elif 'Team Bytes' in val or 'SIH26042' in val:
        r_val.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        r_val.font.bold = True
    else:
        r_val.font.color.rgb = RGBColor(0x21, 0x25, 0x29)

# ==================== SLIDE 2 ====================
slide2 = prs.slides[1]
# Title 1
title2 = [s for s in slide2.shapes if s.name == 'Title 1'][0]
title2.text_frame.clear()
p_t2 = title2.text_frame.paragraphs[0]
r_t2 = p_t2.add_run()
r_t2.text = 'IDEA TITLE: VERNACULAR VOICE'
r_t2.font.name = 'Times New Roman'
r_t2.font.size = Pt(32)
r_t2.font.bold = True
r_t2.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)

# TextBox 8
tb8_s2 = [s for s in slide2.shapes if s.name == 'TextBox 8'][0]
tb8_s2.left = Emu(400000)
tb8_s2.top = Emu(1350000)
tb8_s2.width = Emu(11392000)
tb8_s2.height = Emu(4800000)
tb8_s2.text_frame.clear()
tb8_s2.text_frame.word_wrap = True

def add_header(tf, text):
    p = tf.paragraphs[0] if len(tf.paragraphs) == 1 and not tf.paragraphs[0].text else tf.add_paragraph()
    p.space_before = Pt(2)
    p.space_after = Pt(8)
    r = p.add_run()
    r.text = '❖  ' + text
    r.font.name = 'Arial'
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.underline = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

def add_section_pointer(tf, pointer_title):
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = '•  ' + pointer_title
    r.font.name = 'Arial'
    r.font.size = Pt(15.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

def add_bullet(tf, text, bold_prefix=''):
    p = tf.add_paragraph()
    p.level = 1
    p.space_before = Pt(1.5)
    p.space_after = Pt(2)
    p.line_spacing = 1.12
    
    # Custom bullet character
    r_b = p.add_run()
    r_b.text = '    –  '
    r_b.font.name = 'Arial'
    r_b.font.size = Pt(13.5)
    r_b.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    
    if bold_prefix:
        r_pref = p.add_run()
        r_pref.text = bold_prefix + ': '
        r_pref.font.name = 'Arial'
        r_pref.font.size = Pt(13.5)
        r_pref.font.bold = True
        r_pref.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    r_txt = p.add_run()
    r_txt.text = text
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(13.5)
    r_txt.font.bold = False
    r_txt.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

add_header(tb8_s2.text_frame, 'Proposed Solution (Describe your Idea/Solution/Prototype)')

add_section_pointer(tb8_s2.text_frame, 'Detailed Explanation of the Proposed Solution')
add_bullet(tb8_s2.text_frame, 'AI-powered mobile/web platform translating primary educational content into tribal & regional mother tongues in real time.', 'Platform Overview')
add_bullet(tb8_s2.text_frame, 'Combines speech-to-text (STT), text-to-speech (TTS), and neural machine translation (NMT) for two-way native interaction.', 'Multimodal Core')
add_bullet(tb8_s2.text_frame, 'Generates bilingual teaching materials (synchronized text + audio) with offline-first caching for zero-connectivity classrooms.', 'Offline Pedagogic Flow')

add_section_pointer(tb8_s2.text_frame, 'How It Addresses the Problem')
add_bullet(tb8_s2.text_frame, 'Eliminates learning gaps between standardized state curricula and mother-tongue primary students, directly supporting NEP 2020.', 'Bridging Curriculum Gap')
add_bullet(tb8_s2.text_frame, 'Provides educators with instant automated localization tools, removing dependency on scarce tribal-language human translators.', 'Teacher Empowerment')

add_section_pointer(tb8_s2.text_frame, 'Innovation and Uniqueness of the Solution')
add_bullet(tb8_s2.text_frame, 'Modular plug-and-play language architecture allows rapid addition of new tribal/regional dialects without modifying core engine.', 'Modular Language Plugins')
add_bullet(tb8_s2.text_frame, 'Quantized offline on-device AI inference engineered specifically for ultra-low-cost Android devices, unlike cloud-bound tools.', 'Edge-Optimized Inference')

prs.save('scratch/step1_out.pptx')
print('Step 1 saved successfully')
