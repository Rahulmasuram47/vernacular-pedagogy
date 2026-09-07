import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide1 = prs.slides[0]
for s in slide1.shapes:
    if s.name == 'TextBox 9':
        print('TextBox 9 bounds:', s.left, s.top, s.width, s.height)
        for i, p in enumerate(s.text_frame.paragraphs):
            print(f'P{i}: text={repr(p.text)}')
            for r in p.runs:
                print(f'   run: text={repr(r.text)}, font={r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.font.bold}')
