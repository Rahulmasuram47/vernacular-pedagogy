import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
print('Initial slide count:', len(prs.slides))

# delete slide 7 (index 6)
rId = prs.slides._sldIdLst[6].rId
prs.part.drop_rel(rId)
del prs.slides._sldIdLst[6]

print('After deletion slide count:', len(prs.slides))
prs.save('scratch/test_del.pptx')
print('Saved test_del.pptx')
