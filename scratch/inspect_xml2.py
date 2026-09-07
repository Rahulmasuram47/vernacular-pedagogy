import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide2 = prs.slides[1]
tb8 = [s for s in slide2.shapes if s.name == 'TextBox 8'][0]
print(tb8._element.xml)
