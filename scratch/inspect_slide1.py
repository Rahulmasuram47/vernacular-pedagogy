import pptx

prs = pptx.Presentation(r'C:\Users\harsh\Downloads\SIH2026-IDEA-Presentation-Format.pptx')
slide1 = prs.slides[0]
for s in slide1.shapes:
    print(f'Shape {s.name}: left={s.left}, top={s.top}, width={s.width}, height={s.height}')
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            print('   P:', repr(p.text))
